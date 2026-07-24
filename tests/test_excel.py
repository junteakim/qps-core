from openpyxl import load_workbook

from qps_core.excel import build_review_workbook, verify_formula_workbook
from qps_core.models import CostLine
from qps_core.quotation import QuoteEngine


def test_workbook_keeps_calculated_cells_as_formulas(tmp_path) -> None:
    snapshot = QuoteEngine().calculate(
        quote_id="DEMO-Q-002",
        currency="DEMO",
        lines=[
            CostLine("Material", "2", "1200", "SYNTHETIC"),
            CostLine("Labor", "10", "75", "SYNTHETIC"),
        ],
        overhead_percent="10",
        profit_percent="8",
    )
    path = build_review_workbook(snapshot, tmp_path / "review.xlsx")

    assert verify_formula_workbook(path) == []
    workbook = load_workbook(path, data_only=False)
    assert workbook["Inputs"]["D2"].value == "=B2*C2"
    assert workbook["Summary"]["B9"].value == "=B6+B8"


def test_workbook_escapes_formula_like_text_inputs(tmp_path) -> None:
    snapshot = QuoteEngine().calculate(
        quote_id="+DEMO",
        currency="-DEMO",
        lines=[CostLine("=1+1", "1", "1", "@DEMO")],
    )
    path = build_review_workbook(snapshot, tmp_path / "safe.xlsx")
    workbook = load_workbook(path, data_only=False)

    assert workbook["Inputs"]["A2"].value == "'=1+1"
    assert workbook["Inputs"]["E2"].value == "'@DEMO"
    assert workbook["Summary"]["B2"].value == "'+DEMO"
    assert workbook["Summary"]["B10"].value == "'-DEMO"
    assert all(
        workbook["Inputs"][coordinate].data_type == "s" for coordinate in ("A2", "E2")
    )
