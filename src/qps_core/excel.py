from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from qps_core.models import QuoteSnapshot

FORMULA_ERROR_TOKENS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!")


def _safe_excel_text(value: str) -> str:
    text = str(value)
    if text.startswith(("\t", "\r", "\n")) or text.lstrip().startswith(
        ("=", "+", "-", "@")
    ):
        return "'" + text
    return text


def build_review_workbook(
    snapshot: QuoteSnapshot,
    output_path: str | Path,
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    summary = workbook.create_sheet("Summary")
    contract = workbook.create_sheet("Formula Contract")

    headers = (
        "Category",
        "Quantity",
        "Unit Cost",
        "Amount",
        "Provenance",
        "Cell Role",
    )
    inputs.append(headers)
    for cell in inputs[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    formula_contract: list[tuple[str, str, str]] = []
    for row_index, line in enumerate(snapshot.lines, start=2):
        inputs.cell(row_index, 1, _safe_excel_text(line.category))
        inputs.cell(row_index, 2, float(line.quantity))
        inputs.cell(row_index, 3, float(line.unit_cost))
        inputs.cell(row_index, 4, f"=B{row_index}*C{row_index}")
        inputs.cell(row_index, 5, _safe_excel_text(line.provenance))
        inputs.cell(row_index, 6, "formula-first")
        formula_contract.append(
            ("Inputs", f"D{row_index}", f"=B{row_index}*C{row_index}")
        )

    last_input_row = len(snapshot.lines) + 1
    summary.append(("Synthetic Quote Summary", None))
    summary.append(("Quote ID", _safe_excel_text(snapshot.quote_id)))
    summary.append(("Base Cost", f"=SUM(Inputs!D2:D{last_input_row})"))
    summary.append(("Overhead Percent", float(snapshot.overhead_percent)))
    summary.append(("Overhead Amount", "=B3*B4/100"))
    summary.append(("Subtotal", "=B3+B5"))
    summary.append(("Profit Percent", float(snapshot.profit_percent)))
    summary.append(("Profit Amount", "=B6*B7/100"))
    summary.append(("Grand Total", "=B6+B8"))
    summary.append(("Currency", _safe_excel_text(snapshot.currency)))

    summary["A1"].font = Font(bold=True, size=14)
    formula_contract.extend(
        [
            ("Summary", "B3", f"=SUM(Inputs!D2:D{last_input_row})"),
            ("Summary", "B5", "=B3*B4/100"),
            ("Summary", "B6", "=B3+B5"),
            ("Summary", "B8", "=B6*B7/100"),
            ("Summary", "B9", "=B6+B8"),
        ]
    )

    contract.append(("Sheet", "Cell", "Expected Formula"))
    for row in formula_contract:
        contract.append(row)
    for cell in contract[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")

    inputs.freeze_panes = "A2"
    for sheet in (inputs, summary, contract):
        for column in sheet.columns:
            width = min(
                42,
                max(len(str(cell.value or "")) for cell in column) + 2,
            )
            sheet.column_dimensions[column[0].column_letter].width = width
    for row in range(2, last_input_row + 1):
        inputs.cell(row, 3).number_format = "#,##0.00"
        inputs.cell(row, 4).number_format = "#,##0.00"
    for row in (3, 5, 6, 8, 9):
        summary.cell(row, 2).number_format = "#,##0.00"

    workbook.save(output)
    return output


def verify_formula_workbook(path: str | Path) -> list[str]:
    workbook = load_workbook(path, data_only=False)
    issues: list[str] = []
    required_sheets = {"Inputs", "Summary", "Formula Contract"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        issues.append(f"missing sheets: {sorted(missing)}")
        return issues

    contract = workbook["Formula Contract"]
    for sheet_name, coordinate, expected in contract.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not sheet_name or not coordinate:
            continue
        actual = workbook[str(sheet_name)][str(coordinate)].value
        if actual != expected:
            issues.append(
                f"{sheet_name}!{coordinate}: expected {expected!r}, got {actual!r}"
            )

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for token in FORMULA_ERROR_TOKENS:
                        if token in cell.value:
                            issues.append(
                                f"{sheet.title}!{cell.coordinate}: contains {token}"
                            )
    return issues


def write_snapshot(snapshot: QuoteSnapshot, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(snapshot.to_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return output
